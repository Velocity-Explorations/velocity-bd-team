#!/usr/bin/env python3
"""
Re-ingest March 25 batch into ChromaDB with per-file source tracking.
Deletes generic 'past_submissions' and 'one_pagers' chunks,
re-embeds with actual filenames as source metadata.
"""

import hashlib
import logging
import os
import sys
import time
from pathlib import Path

import requests

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
log = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent.parent
CHROMA_PATH = str(BASE_DIR / "data" / "chroma")
SUMMARY_PATH = BASE_DIR / "data" / "reingestion-summary.md"

PAST_DIR = BASE_DIR / "data" / "bd-history" / "past-submissions" / "Past Submissions"
ONEP_DIR = BASE_DIR / "data" / "bd-history" / "one-pagers" / "One Pagers"

OLLAMA_URL = "http://localhost:11434/api/embeddings"

CHUNK_WORDS = 400
OVERLAP_WORDS = 40
MIN_CHUNK_CHARS = 50

# Files already properly tagged — skip these
SKIP_PREFIXES = [
    "KAIROS_PWS",
    "Echo_C2_AF",
    "Velocity_AI_Policy",
    "Velocity_Capabilities_Brief_2025_FULL",
    "Velocity_Company_Overview",
    "Velocity_Core_Competencies",
    "Velocity_Past_Performance",
    "Velocity_Security_Clearance",
    "Velocity_Capabilities_Brief_2025.pdf",
]

SUPPORTED_EXTS = {".docx", ".pdf", ".pptx", ".xlsx", ".txt"}


# --- Text extraction ---

def extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_pdf(path: Path) -> str:
    import fitz
    doc = fitz.open(str(path))
    text = []
    for page in doc:
        text.append(page.get_text())
    doc.close()
    return "\n".join(text)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        text.append(t)
    return "\n".join(text)


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    text = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None]
            if vals:
                text.append(" | ".join(vals))
    wb.close()
    return "\n".join(text)


def extract_txt(path: Path) -> str:
    return path.read_text(errors="replace")


EXTRACTORS = {
    ".docx": extract_docx,
    ".pdf": extract_pdf,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
    ".txt": extract_txt,
}


# --- Doc type inference ---

def infer_doc_type(filename: str) -> str:
    fn = filename.lower()
    if any(k in fn for k in ["contract", "fa864", "pws"]):
        return "contract"
    if any(k in fn for k in ["proposal", "volume", "aamac"]):
        return "proposal"
    if any(k in fn for k in ["capability", "overview", "brief", "one-pager", "one_pager", "one pager"]):
        return "capability_statement"
    if any(k in fn for k in ["white_paper", "whitepaper", "white paper"]):
        return "white_paper"
    if any(k in fn for k in ["sow", "statement of work", "statement+of+work"]):
        return "sow"
    if "charter" in fn:
        return "charter"
    if any(k in fn for k in ["rfi", "rfq"]):
        return "rfi"
    return "other"


# --- Chunking (word-based) ---

def chunk_text(text: str, chunk_words: int = CHUNK_WORDS, overlap: int = OVERLAP_WORDS) -> list[str]:
    text = text.strip()
    if not text:
        return []
    words = text.split()
    if not words:
        return []

    chunks = []
    start = 0
    while start < len(words):
        end = start + chunk_words
        chunk = " ".join(words[start:end])
        if len(chunk) >= MIN_CHUNK_CHARS:
            chunks.append(chunk)
        start += chunk_words - overlap

    return chunks


# --- Embedding ---

def embed(text: str) -> list[float]:
    r = requests.post(OLLAMA_URL, json={"model": "nomic-embed-text", "prompt": text}, timeout=30)
    r.raise_for_status()
    return r.json()["embedding"]


# --- File collection ---

def should_skip(filename: str) -> bool:
    for prefix in SKIP_PREFIXES:
        if filename.startswith(prefix):
            return True
    return False


def collect_files(directory: Path) -> list[Path]:
    """Collect all processable files, including inside extracted zip directories."""
    files = []
    if not directory.exists():
        return files
    for p in sorted(directory.rglob("*")):
        if not p.is_file():
            continue
        if p.name.startswith(".") or p.name.startswith("~"):
            continue
        if p.suffix.lower() not in SUPPORTED_EXTS:
            continue
        if should_skip(p.name):
            continue
        # Skip .zip files themselves (we process their extracted dirs)
        files.append(p)
    return files


# --- Main ---

def main():
    import chromadb

    client = chromadb.PersistentClient(path=CHROMA_PATH)
    col = client.get_or_create_collection(
        "bd_documents",
        metadata={"description": "Velocity past BD submissions and capability documents"},
    )

    # Step 1: Delete generic-source chunks
    log.info("Deleting chunks with source='past_submissions' or 'one_pagers'...")
    result = col.get(include=["metadatas"])
    delete_ids = []
    for doc_id, meta in zip(result["ids"], result["metadatas"]):
        src = meta.get("source", "")
        if src in ("past_submissions", "one_pagers"):
            delete_ids.append(doc_id)

    if delete_ids:
        # ChromaDB delete in batches
        batch = 500
        for i in range(0, len(delete_ids), batch):
            col.delete(ids=delete_ids[i:i + batch])
        log.info(f"Deleted {len(delete_ids)} generic-source chunks")
    else:
        log.info("No generic-source chunks found to delete")

    remaining = col.count()
    log.info(f"Collection now has {remaining} items (properly tagged)")

    # Step 2: Collect files from both directories
    files_past = collect_files(PAST_DIR)
    files_onep = collect_files(ONEP_DIR)
    all_files = files_past + files_onep
    log.info(f"Found {len(all_files)} files to process ({len(files_past)} past submissions, {len(files_onep)} one-pagers)")

    # Step 3: Process each file
    processed = []
    failed = []
    total_chunks = 0

    for fpath in all_files:
        fname = fpath.name
        ext = fpath.suffix.lower()
        doc_type = infer_doc_type(fname)

        try:
            text = EXTRACTORS[ext](fpath)
            char_count = len(text.strip())
            if char_count == 0:
                log.warning(f"  SKIP (empty): {fname}")
                failed.append((fname, "empty text"))
                continue

            chunks = chunk_text(text)
            if not chunks:
                log.warning(f"  SKIP (no viable chunks): {fname}")
                failed.append((fname, "no viable chunks"))
                continue

            # Embed and upsert each chunk
            chunk_ids = []
            chunk_docs = []
            chunk_embs = []
            chunk_metas = []

            for i, chunk in enumerate(chunks):
                chunk_id = hashlib.md5(f"{fname}:{i}:{chunk[:50]}".encode()).hexdigest()
                emb = embed(chunk)
                meta = {
                    "source": fname,
                    "doc_type": doc_type,
                    "chunk": i,
                }
                chunk_ids.append(chunk_id)
                chunk_docs.append(chunk)
                chunk_embs.append(emb)
                chunk_metas.append(meta)

            # Upsert in batches of 100
            for i in range(0, len(chunk_ids), 100):
                end = min(i + 100, len(chunk_ids))
                col.upsert(
                    ids=chunk_ids[i:end],
                    documents=chunk_docs[i:end],
                    embeddings=chunk_embs[i:end],
                    metadatas=chunk_metas[i:end],
                )

            total_chunks += len(chunks)
            processed.append((fname, doc_type, len(chunks), char_count))
            log.info(f"  {fname} | {char_count} chars | {len(chunks)} chunks | {doc_type}")

        except Exception as e:
            log.error(f"  FAILED {fname}: {e}")
            failed.append((fname, str(e)))

    final_count = col.count()
    log.info(f"\nDone: {len(processed)} files, {total_chunks} new chunks. Collection total: {final_count}")

    # Step 4: Write summary
    summary_lines = [
        "# RAG Re-Ingestion Summary (Per-File Source Tracking)",
        f"\n**Date**: {time.strftime('%Y-%m-%d %H:%M')}",
        f"**Files processed**: {len(processed)}",
        f"**Generic chunks deleted**: {len(delete_ids)}",
        f"**New chunks added**: {total_chunks}",
        f"**Collection total**: {final_count}",
        "\n## Files Processed\n",
        "| File | Doc Type | Chunks | Chars |",
        "|------|----------|--------|-------|",
    ]
    for fname, dtype, nchunks, chars in processed:
        summary_lines.append(f"| {fname} | {dtype} | {nchunks} | {chars} |")

    if failed:
        summary_lines.append("\n## Failed/Skipped Files\n")
        for fname, reason in failed:
            summary_lines.append(f"- **{fname}**: {reason}")

    # Source distribution check
    summary_lines.append("\n## Source Distribution (Post-Ingestion)\n")
    result = col.get(include=["metadatas"])
    sources = {}
    for m in result["metadatas"]:
        s = m.get("source", "unknown")
        sources[s] = sources.get(s, 0) + 1
    for s, c in sorted(sources.items()):
        summary_lines.append(f"- `{s}`: {c} chunks")

    # Sanity check
    summary_lines.append("\n## Sanity Check Query\n")
    summary_lines.append('Query: "SBIR artificial intelligence C2 software"\n')
    try:
        q_emb = embed("SBIR artificial intelligence C2 software")
        results = col.query(query_embeddings=[q_emb], n_results=5)
        for i, meta in enumerate(results["metadatas"][0]):
            dist = results["distances"][0][i] if "distances" in results else "?"
            summary_lines.append(f"{i+1}. **{meta['source']}** (type={meta['doc_type']}, chunk={meta['chunk']}, dist={dist})")
    except Exception as e:
        summary_lines.append(f"Query failed: {e}")

    SUMMARY_PATH.parent.mkdir(parents=True, exist_ok=True)
    SUMMARY_PATH.write_text("\n".join(summary_lines) + "\n")
    log.info(f"Summary written to {SUMMARY_PATH}")

    # Return stats for caller
    print(f"\nRE-INGESTION COMPLETE: {len(processed)} files, {total_chunks} chunks (deleted {len(delete_ids)} generic). Collection: {final_count}")


if __name__ == "__main__":
    main()
