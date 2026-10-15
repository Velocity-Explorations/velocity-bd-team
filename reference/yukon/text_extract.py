"""
Yukon — Shared Text Extraction & Embedding Utilities
Used by ingest_past_submissions.py and ingest_submitted_draft.py.
"""

import logging
import re
from pathlib import Path
from typing import List

import requests

log = logging.getLogger(__name__)

CHUNK_SIZE = 2000
OLLAMA_URL = "http://localhost:11434/api/embeddings"
EMBED_MODEL = "nomic-embed-text"


def extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_pdf(path: Path) -> str:
    import fitz
    doc = fitz.open(str(path))
    text = []
    for page in doc:
        text.append(page.get_text())
    doc.close()
    return "\n".join(text)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        text.append(t)
    return "\n".join(text)


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    text = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None]
            if vals:
                text.append(" | ".join(vals))
    wb.close()
    return "\n".join(text)


EXTRACTORS = {
    ".docx": extract_docx,
    ".pdf": extract_pdf,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
}


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    extractor = EXTRACTORS.get(ext)
    if not extractor:
        return ""
    return extractor(path)


def infer_doc_type(filename: str) -> str:
    fn = filename.lower()
    if any(k in fn for k in ["proposal", "tech volume", "vol1", "vol2"]):
        return "proposal"
    if any(k in fn for k in ["requirement", "analysis"]):
        return "requirements"
    if any(k in fn for k in ["outline"]):
        return "outline"
    if any(k in fn for k in ["draft", "response"]):
        return "draft_response"
    if any(k in fn for k in ["checklist", "submission"]):
        return "checklist"
    if any(k in fn for k in ["sow", "pws", "statement of work"]):
        return "sow"
    if any(k in fn for k in ["white paper", "whitepaper"]):
        return "white_paper"
    if "charter" in fn:
        return "charter"
    if "rfi" in fn:
        return "rfi"
    if any(k in fn for k in ["quad chart", "one pager", "one-page", "capability"]):
        return "capability_statement"
    return "other"


def make_slug(filename: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", filename.rsplit(".", 1)[0])
    return slug.strip("_").lower()


def chunk_text(text: str, size: int = CHUNK_SIZE) -> List[str]:
    text = text.strip()
    if not text:
        return []
    chunks = []
    for i in range(0, len(text), size):
        chunk = text[i:i + size].strip()
        if chunk:
            chunks.append(chunk)
    return chunks


def embed(text: str) -> List[float]:
    r = requests.post(OLLAMA_URL, json={"model": EMBED_MODEL, "prompt": text}, timeout=60)
    r.raise_for_status()
    return r.json()["embedding"]
