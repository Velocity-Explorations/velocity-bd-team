#!/usr/bin/env python3
"""
Ingest past BD submissions into ChromaDB.
Extracts text from .docx, .pdf, .pptx, .xlsx files,
chunks them, embeds via Ollama nomic-embed-text, and upserts into ChromaDB.
"""

import json
import logging
import os
import re
import sys
import time
from pathlib import Path

import requests
from dotenv import load_dotenv

load_dotenv(Path(__file__).parent.parent / "configs" / "velocity" / ".env")

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
log = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data" / "bd-history" / "past-submissions" / "Past Submissions"
CHROMA_PATH = str(BASE_DIR / "data" / "chroma")
SUMMARY_PATH = BASE_DIR / "data" / "bd-history" / "ingestion-summary.md"

CHUNK_SIZE = 2000  # ~500 tokens
OLLAMA_URL = "http://localhost:11434/api/embeddings"
DISCORD_API = "https://discord.com/api/v10"
DISCORD_CHANNEL = "1483788952640159824"


# --- Text extraction ---

def extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_pdf(path: Path) -> str:
    import fitz
    doc = fitz.open(str(path))
    text = []
    for page in doc:
        text.append(page.get_text())
    doc.close()
    return "\n".join(text)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        text.append(t)
    return "\n".join(text)


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    text = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None]
            if vals:
                text.append(" | ".join(vals))
    wb.close()
    return "\n".join(text)


EXTRACTORS = {
    ".docx": extract_docx,
    ".pdf": extract_pdf,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
}


# --- Doc type inference ---

def infer_doc_type(filename: str) -> str:
    fn = filename.lower()
    if any(k in fn for k in ["proposal", "tech volume", "vol1", "vol2"]):
        return "proposal"
    if any(k in fn for k in ["sow", "pws", "statement of work"]):
        return "sow"
    if any(k in fn for k in ["contract", "fa864924"]):
        return "contract"
    if any(k in fn for k in ["white paper", "whitepaper", "white_paper"]):
        return "white_paper"
    if "charter" in fn:
        return "charter"
    if "rfi" in fn:
        return "rfi"
    if any(k in fn for k in ["quad chart", "one pager", "one-page", "one_page"]):
        return "capability_statement"
    if any(k in fn for k in ["los", "pitch", "cos"]):
        return "capability_statement"
    return "other"


# --- Chunking ---

def make_slug(filename: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", filename.rsplit(".", 1)[0])
    return slug.strip("_").lower()


def chunk_text(text: str, size: int = CHUNK_SIZE) -> list[str]:
    text = text.strip()
    if not text:
        return []
    chunks = []
    for i in range(0, len(text), size):
        chunk = text[i:i + size].strip()
        if chunk:
            chunks.append(chunk)
    return chunks


# --- Embedding ---

def embed(text: str) -> list[float]:
    r = requests.post(OLLAMA_URL, json={"model": "nomic-embed-text", "prompt": text}, timeout=60)
    r.raise_for_status()
    return r.json()["embedding"]


# --- Discord ---

def _get_bot_token() -> str:
    token = os.environ.get("DISCORD_BOT_TOKEN")
    if not token:
        raise RuntimeError("DISCORD_BOT_TOKEN not set in environment")
    return token


def post_discord(message: str):
    token = _get_bot_token()
    headers = {"Authorization": f"Bot {token}", "Content-Type": "application/json"}
    url = f"{DISCORD_API}/channels/{DISCORD_CHANNEL}/messages"
    resp = requests.post(url, headers=headers, json={"content": message}, timeout=10)
    if resp.status_code == 200:
        log.info("Discord message posted.")
    else:
        log.error(f"Discord error {resp.status_code}: {resp.text[:200]}")


# --- Main ---

def main():
    import chromadb

    # Collect files
    files = []
    for p in sorted(DATA_DIR.rglob("*")):
        if p.is_file() and p.suffix.lower() in EXTRACTORS and not p.name.startswith("."):
            files.append(p)

    log.info(f"Found {len(files)} files to process")

    client = chromadb.PersistentClient(path=CHROMA_PATH)
    collection = client.get_or_create_collection(
        "bd_documents",
        metadata={"description": "Velocity past BD submissions and capability documents"},
    )

    processed = []
    failed = []
    all_ids = []
    all_docs = []
    all_embeddings = []
    all_metadatas = []

    for fpath in files:
        fname = fpath.name
        ext = fpath.suffix.lower()
        doc_type = infer_doc_type(fname)
        log.info(f"Processing: {fname} ({doc_type})")

        try:
            text = EXTRACTORS[ext](fpath)
            if not text.strip():
                log.warning(f"  Empty text: {fname}")
                failed.append((fname, "empty text"))
                continue

            chunks = chunk_text(text)
            slug = make_slug(fname)

            for i, chunk in enumerate(chunks):
                doc_id = f"{slug}_{i}"
                emb = embed(chunk)
                meta = {
                    "filename": fname,
                    "doc_type": doc_type,
                    "source": "past_submissions",
                    "chunk_index": i,
                    "total_chunks": len(chunks),
                }
                all_ids.append(doc_id)
                all_docs.append(chunk)
                all_embeddings.append(emb)
                all_metadatas.append(meta)

            processed.append((fname, doc_type, len(chunks)))
            log.info(f"  -> {len(chunks)} chunks")

        except Exception as e:
            log.error(f"  FAILED: {e}")
            failed.append((fname, str(e)))

    # Upsert in batches of 100
    total_chunks = len(all_ids)
    log.info(f"Upserting {total_chunks} chunks into ChromaDB...")
    batch_size = 100
    for i in range(0, total_chunks, batch_size):
        end = min(i + batch_size, total_chunks)
        collection.upsert(
            ids=all_ids[i:end],
            documents=all_docs[i:end],
            embeddings=all_embeddings[i:end],
            metadatas=all_metadatas[i:end],
        )
    log.info(f"ChromaDB collection 'bd_documents' now has {collection.count()} items")

    # Write summary
    summary_lines = [
        "# BD Past Submissions Ingestion Summary",
        f"\n**Date**: {time.strftime('%Y-%m-%d %H:%M')}",
        f"**Files processed**: {len(processed)}",
        f"**Total chunks embedded**: {total_chunks}",
        f"**Collection size**: {collection.count()}",
        "\n## Files Processed\n",
        "| File | Doc Type | Chunks |",
        "|------|----------|--------|",
    ]
    for fname, dtype, nchunks in processed:
        summary_lines.append(f"| {fname} | {dtype} | {nchunks} |")

    if failed:
        summary_lines.append("\n## Failed Files\n")
        for fname, reason in failed:
            summary_lines.append(f"- **{fname}**: {reason}")

    # Sanity check query
    summary_lines.append("\n## Sanity Check Query\n")
    summary_lines.append('Query: "SBIR artificial intelligence C2 software USAF"\n')
    try:
        q_emb = embed("SBIR artificial intelligence C2 software USAF")
        results = collection.query(query_embeddings=[q_emb], n_results=5)
        top_matches = []
        for i, (doc_id, meta) in enumerate(zip(results["ids"][0], results["metadatas"][0])):
            summary_lines.append(f"{i+1}. **{meta['filename']}** (type={meta['doc_type']}, chunk={meta['chunk_index']})")
            top_matches.append(meta["filename"])
    except Exception as e:
        summary_lines.append(f"Query failed: {e}")
        top_matches = []

    SUMMARY_PATH.parent.mkdir(parents=True, exist_ok=True)
    SUMMARY_PATH.write_text("\n".join(summary_lines) + "\n")
    log.info(f"Summary written to {SUMMARY_PATH}")

    # Post to Discord
    top3 = list(dict.fromkeys(top_matches))[:3]  # unique filenames, top 3
    top3_str = ", ".join(top3) if top3 else "N/A"
    discord_msg = (
        f"\u2705 Past Submissions embedded \u2014 {len(processed)} files, "
        f"{total_chunks} chunks in ChromaDB. Ready for Step 3 (One Pagers zip). "
        f"Sample matches: {top3_str}"
    )
    try:
        post_discord(discord_msg)
    except Exception as e:
        log.error(f"Discord post failed: {e}")

    print(f"\nDone: {len(processed)} files, {total_chunks} chunks embedded.")


if __name__ == "__main__":
    main()
