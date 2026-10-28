"""Word files for the workspace's documents (PRO-579): structure, inline styles, placeholders, page breaks and format."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt

from drafting import word

SAMPLE = """# Edge AI white paper

**Company:** Velocity Explorations | CAGE 9KCQ5

---

## 1. Technical Approach

Echo C2 runs **offline** on *rugged* laptops; see `echo-c2` and [the notice](https://sam.gov/opp/1).
It explains each recommendation.

[NEEDS INPUT: latency figures from the Cobra Gold report]

- Fuses sensor reports
  - and spot reports
- Recommends courses of action

1. Outline approved
2. Draft complete

| Item | Owner | Due |
|---|---|---|
| ☐ Bios | David Wood | 2026-11-20 |

> Evaluators read the opening first.
"""


def _texts(doc: Document) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]


def test_a_document_keeps_its_headings_lists_tables_and_page_break(tmp_path: Path) -> None:
    path = word.to_docx(SAMPLE, tmp_path / "draft_v1.docx", page_breaks=True)
    doc = Document(str(path))
    texts = _texts(doc)
    assert texts[0] == ("Heading 1", "Edge AI white paper")
    assert ("Heading 2", "1. Technical Approach") in texts
    assert ("Normal", "Echo C2 runs offline on rugged laptops; see echo-c2 and the notice (https://sam.gov/opp/1). "
                      "It explains each recommendation.") in texts
    assert [t for s, t in texts if s.startswith("List Bullet")] == ["Fuses sensor reports", "and spot reports",
                                                                    "Recommends courses of action"]
    assert [s for s, t in texts if t == "and spot reports"] == ["List Bullet 2"]
    assert [t for s, t in texts if s == "List Number"] == ["Outline approved", "Draft complete"]
    assert ("Quote", "Evaluators read the opening first.") in texts
    [table] = doc.tables
    assert [c.text for c in table.rows[0].cells] == ["Item", "Owner", "Due"]
    assert [c.text for c in table.rows[1].cells] == ["☐ Bios", "David Wood", "2026-11-20"]
    assert any('w:br w:type="page"' in p._p.xml for p in doc.paragraphs)


def test_bold_italic_code_and_placeholders_are_runs(tmp_path: Path) -> None:
    doc = Document(str(word.to_docx(SAMPLE, tmp_path / "a.docx")))
    runs = [r for p in doc.paragraphs for r in p.runs]
    assert next(r for r in runs if r.text == "offline").bold
    assert next(r for r in runs if r.text == "rugged").italic
    assert next(r for r in runs if r.text == "echo-c2").font.name == "Consolas"
    placeholder = next(r for r in runs if r.text.startswith("[NEEDS INPUT"))
    assert placeholder.font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert next(r for r in runs if r.text == "Company:").bold


def test_rules_are_page_breaks_only_in_drafts(tmp_path: Path) -> None:
    doc = Document(str(word.to_docx(SAMPLE, tmp_path / "outline.docx")))
    assert not any('w:type="page"' in p._p.xml for p in doc.paragraphs)


def test_the_solicitations_format_is_applied_and_odd_values_keep_the_default(tmp_path: Path) -> None:
    fmt = word.format_from({"font": "Arial", "size": 11, "margins": 0.75, "spacing": 1.15})
    doc = Document(str(word.to_docx(SAMPLE, tmp_path / "b.docx", fmt)))
    normal = doc.styles["Normal"]
    assert (normal.font.name, normal.font.size, normal.paragraph_format.line_spacing) == ("Arial", Pt(11), 1.15)
    assert doc.sections[0].left_margin == Inches(0.75) and doc.sections[0].top_margin == Inches(0.75)
    assert doc.styles["Heading 1"].font.name == "Arial"
    assert word.format_from({"font": "", "size": 40, "margins": "wide", "spacing": None}) == word.DEFAULT
    assert word.format_from(None) == word.DEFAULT
