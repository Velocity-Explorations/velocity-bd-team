"""The Qualify steps (PW6 to PW10, PRO-577): Discovery, Opportunity summary, Alignment, Requirements and
checklist, and the Go or no-go brief, each run with a stand-in model."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from typing import Any

import pytest

from drafting import inputs, qualify, steps
from drafting.prompts import Prompt
from ledger import queries as lq

FENCE = "`" * 3

DISCOVERY = f"""# Discovery summary

The call asks for an edge AI white paper. The attachments are not in the ledger.

{FENCE}json
{{"key_points": ["Due 1 Dec 2026", "Attachments must be fetched"],
  "tasks": {{"call_page": "for_owner", "attachments": "for_owner", "parent_instructions": "done", "key_facts": "done",
            "past_work": "done", "clashes": "done", "record": "done", "summary": "done"}},
  "owner_fetch": ["The call page on SAM.gov, for amendments"],
  "result_expected": {{"date": "2027-02-15", "words": "Selections are notified within 75 days", "where": "Section 6.2"}}}}
{FENCE}
"""
SUMMARY = f"# Opportunity summary\n\nArmy, edge AI, white paper.\n\n{FENCE}json\n{{\"key_points\": [\"Army edge AI\"], \"worth_it\": \"yes\"}}\n{FENCE}"
ALIGNMENT = (f"# Alignment\n\nStrong on C2, a gap on Army past performance.\n\n{FENCE}json\n"
             '{"key_points": ["Strong on C2"], "verdict": "Strong fit on C2, one gap", '
             '"angle": "Lead with Echo C2 at Cobra Gold; team for Army past performance", '
             '"comparison": "Both lead with Echo C2; the angle adds a teaming partner"}\n' + FENCE)
REQUIREMENTS = "# Requirements analysis\n\n## 6. SOLICITATION TERMINOLOGY\nKEYWORDS: edge AI"
CHECKLIST = json.dumps({
    "key_points": ["Ten pages, PDF", "Submit by 30 Nov"],
    "checklist": [
        {"part": "registrations", "text": "SAM.gov registration active, UEI and CAGE", "source": "per notice", "owner": "David Wood",
         "due": "2026-11-20", "done": True},
        {"part": "documents", "text": "White paper: 10 pages max, PDF", "source": "per notice section 3", "owner": "Claude",
         "due": "not a date", "done": False},
        {"part": "snacks", "text": "Order lunch", "owner": "David Wood", "due": None, "done": False},
        {"part": "inputs", "text": "Key personnel bios", "source": None, "owner": "", "due": None, "done": False},
    ],
    "submit_by": "2026-12-05",
    "format": {"font": "Arial", "size": 11, "margins": 1, "spacing": None},
    "milestones": {"outline": "2026-11-10", "draft": "2026-11-18", "final": "2026-12-04", "lunch": "2026-11-01"},
    "questions": [{"question": "Is a cost volume required?", "deadline": "2026-11-15", "channel": "the contracting officer by email"}],
    "sources": [{"name": "Notice W1-26-R-0001", "what": "the call", "version": "original"}],
})
GO_NO_GO = (f"# Go or no-go brief\n\nModerate fit.\n\n{FENCE}json\n"
            '{"key_points": ["Moderate fit"], "fit": "Moderate fit", "recommendation": "go_with_conditions", '
            '"conditions": ["Sign a teaming partner with Army past performance"]}\n' + FENCE)
REPLIES = {"discovery": DISCOVERY, "opportunity summary": SUMMARY, "alignment": ALIGNMENT,
           "analyzing requirements": REQUIREMENTS, "submission checklist": CHECKLIST, "go or no-go brief": GO_NO_GO}


class Model:
    def __init__(self, replies: dict[str, str] | None = None) -> None:
        self.asked: list[tuple[str, Prompt, int]] = []
        self.replies = replies or REPLIES

    def __call__(self, prompt: Prompt, label: str, max_tokens: int) -> str:
        self.asked.append((label, prompt, max_tokens))
        return self.replies[label]

    def prompt(self, label: str) -> Prompt:
        return next(p for asked, p, _ in self.asked if asked == label)


def _context(tmp_path: Path, model: Model, step: str = "discovery", direction: str | None = "Lead with Echo C2",
             note: str | None = None, previous: dict[str, Any] | None = None) -> tuple[steps.StepContext, list[dict[str, str]]]:
    run_id = lq.start_run("sam_gov")
    lq.upsert_call({"id": "sam-1", "portal": "SAM.gov", "title": "Edge AI for C2", "type": "Sources Sought",
                    "department": "DEPT OF THE ARMY", "due_date": "2026-12-01", "url": "https://sam.gov/opp/1",
                    "description": "The Army seeks edge AI."}, run_id)
    lq.set_human("sam-1", "pursue", "David Wood")
    pursuit_id = lq.start_pursuit("sam-1", "David Wood", format="White paper", direction=direction)
    row = lq.claim_next_step()
    progress: list[dict[str, str]] = []
    ctx = steps.StepContext(pursuit=lq.get_pursuit(pursuit_id), step={**row, "step": step, "request_note": note},
                            record=inputs.gather(pursuit_id), work_dir=tmp_path / f"pursuit-{pursuit_id}", ask=model,
                            progress=lambda states: progress.append(dict(states)), previous=previous, today=date(2026, 10, 28))
    return ctx, progress


def test_a_reply_splits_into_its_document_and_its_data() -> None:
    assert steps.split_reply(f"# Doc\ntext\n{FENCE}json\n{{\"a\": 1}}\n{FENCE}\n") == ("# Doc\ntext", {"a": 1})
    document, data = steps.split_reply(f"# Doc\n{FENCE}json\n{{bad\n{FENCE}")
    assert document.startswith("# Doc") and data == {}
    assert steps.split_reply('{"items": []}') == ("", {"items": []})
    assert steps.split_reply("# Only prose") == ("# Only prose", {})


def test_each_step_follows_its_part_of_the_skill() -> None:
    skill = qualify.SKILL.read_text(encoding="utf-8")
    for step, phase in (("discovery", 1), ("summary", 2), ("alignment", 3), ("requirements", 4), ("go_no_go", 5)):
        assert qualify.skill_section(step).startswith(f"## Phase {phase}:")
        assert qualify.QUESTIONS[step] in skill, step
    assert "Section 2, submission checklist" in qualify.checklist_instructions()


def test_discovery_reports_its_tasks_and_when_the_result_is_expected(ledger_db: str, tmp_path: Path) -> None:
    model = Model()
    ctx, progress = _context(tmp_path, model)
    result = qualify.discovery(ctx)
    saved = (ctx.work_dir / "discovery.md").read_text(encoding="utf-8")
    assert saved.startswith("# Discovery summary") and "json" not in saved
    assert result.files == [{"name": "Discovery summary", "local": f"{ctx.work_dir.name}/discovery.md",
                             "word": f"{ctx.work_dir.name}/discovery.docx"}]
    assert (ctx.work_dir / "discovery.docx").exists()
    assert result.summary == "- Due 1 Dec 2026\n- Attachments must be fetched"
    assert result.content["tasks"]["call_page"] == "for_owner" and result.content["tasks"]["summary"] == "done"
    assert result.content["owner_fetch"] == ["The call page on SAM.gov, for amendments"]
    assert result.content["question"] == qualify.QUESTIONS["discovery"]
    assert progress[0] == {task: "working" for task in qualify.TASKS} and progress[-1] == result.content["tasks"]
    pursuit = lq.get_pursuit(ctx.pursuit["id"])
    assert (pursuit["result_expected_by"], pursuit["result_expected_source"]) == (
        "2027-02-15", '"Selections are notified within 75 days" (Section 6.2)')
    system, user = model.prompt("discovery")
    assert "## Phase 1: Discovery" in system and "cannot open web pages" in system
    assert "Edge AI for C2" in user and "CLASH CHECK" in user


def test_without_a_date_in_the_call_the_result_is_expected_90_days_after_the_deadline(ledger_db: str, tmp_path: Path) -> None:
    reply = DISCOVERY.replace('"date": "2027-02-15"', '"date": null')
    ctx, _ = _context(tmp_path, Model({**REPLIES, "discovery": reply}))
    qualify.discovery(ctx)
    pursuit = lq.get_pursuit(ctx.pursuit["id"])
    assert (pursuit["result_expected_by"], pursuit["result_expected_source"]) == ("2027-03-01", "default, 90 days after the deadline")


def test_unknown_task_states_are_left_for_the_owner(ledger_db: str, tmp_path: Path) -> None:
    reply = DISCOVERY.replace('"call_page": "for_owner"', '"call_page": "maybe later"').replace('"summary": "done"', '"summary": "x"')
    ctx, _ = _context(tmp_path, Model({**REPLIES, "discovery": reply}))
    tasks = qualify.discovery(ctx).content["tasks"]
    assert (tasks["call_page"], tasks["summary"]) == ("for_owner", "done")


def test_the_opportunity_summary_reads_discovery(ledger_db: str, tmp_path: Path) -> None:
    model = Model()
    ctx, _ = _context(tmp_path, model, "summary")
    ctx.save("discovery.md", "# Discovery summary\nFound the BAA.", "Discovery summary")
    result = qualify.summary(ctx)
    assert (ctx.work_dir / "opportunity_summary.md").exists()
    assert result.content["worth_it"] == "yes" and result.summary == "- Army edge AI"
    assert "Found the BAA." in model.prompt("opportunity summary").user


def test_alignment_sets_the_angle_beside_the_owners_direction(ledger_db: str, tmp_path: Path) -> None:
    model = Model()
    ctx, _ = _context(tmp_path, model, "alignment")
    result = qualify.alignment(ctx)
    assert result.content["verdict"] == "Strong fit on C2, one gap"
    assert result.content["angle"].startswith("Lead with Echo C2 at Cobra Gold")
    assert result.content["direction"] == "Lead with Echo C2"
    assert "## Phase 3: Alignment" in model.prompt("alignment").system


def test_requirements_save_the_analysis_the_checklist_and_the_dates(ledger_db: str, tmp_path: Path) -> None:
    model = Model()
    ctx, _ = _context(tmp_path, model, "requirements")
    result = qualify.requirements(ctx)
    assert (ctx.work_dir / "requirements.md").read_text(encoding="utf-8").startswith("# Requirements analysis")
    rendered = (ctx.work_dir / "submission_checklist.md").read_text(encoding="utf-8")
    assert "2.1 Registrations and eligibility" in rendered and "☑" in rendered and "☐" in rendered
    assert "Is a cost volume required?" in rendered and "Notice W1-26-R-0001" in rendered

    items = lq.checklist(ctx.pursuit["id"])
    assert [(i["part"], i["text"], i["owner"], i["due_date"], i["done"]) for i in items] == [
        ("registrations", "SAM.gov registration active, UEI and CAGE", "David Wood", "2026-11-20", 1),
        ("documents", "White paper: 10 pages max, PDF", "Claude", None, 0),
        ("inputs", "Key personnel bios", "David Wood", None, 0),
    ]
    # The model's submit-by fell after the deadline, so it is one business day before; milestones after it go.
    dates = json.loads(lq.get_pursuit(ctx.pursuit["id"])["target_dates"])
    assert dates == {"submit_by": "2026-11-30", "outline": "2026-11-10", "draft": "2026-11-18"}
    assert result.content["submit_by"] == "2026-11-30" and result.content["checklist_items"] == 3
    assert result.content["format"] == {"font": "Arial", "size": 11.0, "margins": 1.0, "spacing": 1.0}
    from docx import Document
    assert Document(str(ctx.work_dir / "requirements.docx")).styles["Normal"].font.name == "Arial"
    assert Document(str(ctx.work_dir / "submission_checklist.docx")).tables
    assert [f["name"] for f in result.files] == ["Requirements analysis", "Requirements and submission checklist"]
    assert "Section 2, submission checklist" in model.prompt("submission checklist").system
    assert "KEYWORDS: edge AI" in model.prompt("submission checklist").user


def test_the_go_or_no_go_brief_recommends_and_names_its_conditions(ledger_db: str, tmp_path: Path) -> None:
    model = Model()
    ctx, _ = _context(tmp_path, model, "go_no_go")
    result = qualify.go_no_go(ctx)
    assert (result.content["fit"], result.content["recommendation"]) == ("Moderate fit", "go_with_conditions")
    assert result.content["conditions"] == ["Sign a teaming partner with Army past performance"]
    ctx.ask = Model({**REPLIES, "go or no-go brief": GO_NO_GO.replace(
        '"Moderate fit", "recommendation": "go_with_conditions"', '"Great", "recommendation": "maybe"')})
    again = qualify.go_no_go(ctx).content
    assert (again["fit"], again["recommendation"]) == (None, None)


def test_asking_for_changes_puts_the_note_and_the_last_version_in_the_prompt(ledger_db: str, tmp_path: Path) -> None:
    model = Model()
    previous = {"content": json.dumps({"document": "# Opportunity summary\nThe first version."})}
    ctx, _ = _context(tmp_path, model, "summary", note="Say more about the funding", previous=previous)
    qualify.summary(ctx)
    user = model.prompt("opportunity summary").user
    assert "THE OWNER ASKED FOR CHANGES TO YOUR LAST VERSION\nSay more about the funding" in user
    assert "The first version." in user


@pytest.mark.parametrize("step", list(qualify.RUNNERS))
def test_no_prompt_or_document_carries_a_dash(ledger_db: str, tmp_path: Path, step: str) -> None:
    model = Model()
    ctx, _ = _context(tmp_path, model, step)
    qualify.RUNNERS[step](ctx)
    for _, prompt, _ in model.asked:
        for text in (prompt.system, prompt.user):
            assert chr(0x2013) not in text and chr(0x2014) not in text
