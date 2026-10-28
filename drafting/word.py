"""Word files for the proposal workspace (PRO-579): each Markdown document the worker saves gets a .docx
beside it, in the format the solicitation sets.

    to_docx(markdown, path, Format(font="Arial", size=11, margins=1.0, spacing=1.0), page_breaks=True)

Markdown stays the working format; the Word file is what people open and what goes to Drive. The
format comes from the Requirements and checklist step (its `format`), with plain defaults (Times New
Roman 12, one-inch margins, single spacing) where the solicitation says nothing. The converter knows
the Markdown the steps write: headings, paragraphs, bullet and numbered lists, tables, quotes, code
blocks, bold, italics, inline code and links. A rule on its own line is a page break in a draft, as
the reference's formatting pass leaves them, and nothing elsewhere, where it only divides sections.
Placeholders a person must fill ([PERSON WRITES: ...], [NEEDS INPUT: ...], [NEEDS EVIDENCE: ...]) are highlighted
yellow, as the proposal-drafting skill asks.
"""
from __future__ import annotations

import re
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
TABLE_DIVIDER = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|\*(?!\s).+?\*|`[^`]+`|\[(?:PERSON WRITES|NEEDS [A-Z]+)\b[^\]]*\]|\[[^\]]+\]\([^)\s]+\))")
PLACEHOLDER = re.compile(r"^\[(?:PERSON WRITES|NEEDS [A-Z]+)\b", re.I)
LINK = re.compile(r"^\[([^\]]+)\]\(([^)\s]+)\)$")


@dataclass(frozen=True)
class Format:
    font: str = "Times New Roman"
    size: float = 12.0
    margins: float = 1.0     # inches, every side
    spacing: float = 1.0     # line spacing, a multiple


DEFAULT = Format()


def format_from(found: Mapping[str, Any] | None) -> Format:
    """The format a solicitation sets, from the Requirements step's `format`; anything missing or out of a
    sensible range keeps the default."""
    found = found if isinstance(found, Mapping) else {}

    def number(key: str, low: float, high: float, default: float) -> float:
        try:
            value = float(found.get(key))
        except (TypeError, ValueError):
            return default
        return value if low <= value <= high else default

    font = str(found.get("font") or "").strip()
    return Format(font=font if 0 < len(font) <= 40 else DEFAULT.font, size=number("size", 8, 16, DEFAULT.size),
                  margins=number("margins", 0.5, 2.0, DEFAULT.margins), spacing=number("spacing", 1.0, 3.0, DEFAULT.spacing))


def _font(style: Any, fmt: Format, size: float | None = None) -> None:
    style.font.name = fmt.font
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), fmt.font)
    if size:
        style.font.size = Pt(size)


def _runs(paragraph: Any, text: str) -> None:
    """Text with its bold, italics, inline code, links and highlighted placeholders, as runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith(("**", "__")) and part.endswith(part[:2]) and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif PLACEHOLDER.match(part):
            paragraph.add_run(part).font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        elif LINK.match(part):
            label, url = LINK.match(part).groups()
            paragraph.add_run(label if label == url else f"{label} ({url})")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _table(doc: Any, lines: list[str]) -> None:
    rows = [_cells(line) for line in lines if not TABLE_DIVIDER.match(line)]
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            paragraph = table.cell(r, c).paragraphs[0]
            _runs(paragraph, row[c] if c < len(row) else "")
            if r == 0:
                for run in paragraph.runs:
                    run.bold = True


def to_docx(markdown: str, path: Path, fmt: Format = DEFAULT, page_breaks: bool = False) -> Path:
    """Write a Markdown document as a Word file in the given format; returns its path."""
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(fmt.margins)
    normal = doc.styles["Normal"]
    _font(normal, fmt, fmt.size)
    normal.paragraph_format.line_spacing = fmt.spacing
    for level in range(1, 7):
        _font(doc.styles[f"Heading {level}"], fmt)
    lines = markdown.replace("\r\n", "\n").split("\n")
    i, prose = 0, []

    def flush() -> None:
        if prose:
            _runs(doc.add_paragraph(), " ".join(prose))
            prose.clear()

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            flush()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                doc.add_paragraph().add_run(lines[i]).font.name = "Consolas"
                i += 1
        elif not line.strip():
            flush()
        elif HEADING.match(line):
            flush()
            hashes, text = HEADING.match(line).groups()
            _runs(doc.add_heading(level=len(hashes)), text)
        elif RULE.match(line):
            flush()
            if page_breaks:
                doc.add_page_break()
        elif line.lstrip().startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            _table(doc, block)
            continue
        elif BULLET.match(line) or NUMBERED.match(line):
            flush()
            bullet = BULLET.match(line)
            indent, text = (bullet or NUMBERED.match(line)).groups()
            style = "List Bullet" if bullet else "List Number"
            deeper = " 2" if len(indent.replace("\t", "    ")) >= 2 else ""
            _runs(doc.add_paragraph(style=style + deeper), text)
        elif line.lstrip().startswith(">"):
            flush()
            _runs(doc.add_paragraph(style="Quote"), line.lstrip()[1:].strip())
        else:
            prose.append(line.strip())
        i += 1
    flush()
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path

